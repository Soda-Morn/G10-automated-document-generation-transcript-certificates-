import tkinter as tk
from tkinter import messagebox
import os
from docxtpl import DocxTemplate
from docx2pdf import convert
from PIL import Image, ImageDraw, ImageFont
import openpyxl
import pandas as pd
from datetime import datetime


# Function for generating certificates
def generate_certificates(excel_file, template_file, output_folder, font_path="arialbd.ttf", font_size=100):
    data = pd.read_excel(excel_file)
    if not os.path.exists(output_folder):
        os.makedirs(output_folder)
    font_name = ImageFont.truetype(font_path, font_size)
    for index, row in data.iterrows():
        name = row["Name"]
        certificate = Image.open(template_file)
        draw = ImageDraw.Draw(certificate)
        if len(name) >= 15 and len(name) < 25:
            name_position = (550, 600)
        elif len(name) >= 10 and len(name) < 15:
            name_position = (700, 600)
        else:
            name_position = (730, 600)
        draw.text(name_position, name, fill="orange", font=font_name)
        output_path = os.path.join(output_folder, "certificate_" + name + ".png")
        certificate.save(output_path)
        print(f"Certificate generated for {name} and saved to {output_path}")
    print("All certificates have been generated!")

# Functions for generating Associate Degree documents
def AssociateExcel_data(filename):
    workbook = openpyxl.load_workbook(filename)
    sheet = workbook.active
    return list(sheet.values)

def AssociateDocument(template, output_directory, student):
    doc = DocxTemplate(template)
    current_date = datetime.now().strftime("%B %d, %Y")
    doc.render({
        'name_kh': student[2],
        'g1': student[4],
        'id_kh': student[0],
        'name_e': student[3],
        'g2': student[5],
        'id_e': student[1],
        'dob_kh': student[6],
        'pro_kh': student[8],
        'dob_e': student[7],
        'pro_e': student[9],
        'ed_kh': student[10],
        'ed_e': student[11],
        'cur_date': current_date
    })
    doc_name = os.path.join(output_directory, f"{student[3]}.docx")
    doc.save(doc_name)
    return doc_name

def AssociateConvertPDF(doc_path, pdf_directory):
    pdf_path = os.path.join(pdf_directory, os.path.splitext(os.path.basename(doc_path))[0] + ".pdf")
    convert(doc_path, pdf_path)
    return pdf_path

def GeneratAssociate(option):
    excel_file = "Associate.xlsx"
    template_file = "WEP Temporary Certificate - Template.docx"
    docx_directory = "Associate_Documents"
    pdf_directory = "Associate_PDF"

    os.makedirs(docx_directory, exist_ok=True)
    os.makedirs(pdf_directory, exist_ok=True)
    data_rows = AssociateExcel_data(excel_file)

    for row in data_rows[1:]:
        if option in ["doc", "both"]:
            doc_path = AssociateDocument(template_file, docx_directory, row)
        if option in ["pdf", "both"]:
            if option == "pdf":
                doc_path = AssociateDocument(template_file, pdf_directory, row)
            AssociateConvertPDF(doc_path, pdf_directory)
            if option == "pdf":
                os.remove(doc_path)
    print(f"All files for option '{option}' have been generated!")

# Functions for generating Transcripts
def TranscriptExcel_data(filename):
    workbook = openpyxl.load_workbook(filename)
    sheet = workbook.active
    return list(sheet.values)

def TranscriptDocument(template, output_directory, row_data):
    doc = DocxTemplate(template)
    current_date = datetime.now().strftime("%B %d, %Y")
    doc.render({
        "student_id": row_data[0],
        "first_name": row_data[1],
        "last_name": row_data[2],
        "logic": row_data[3],
        "l_g": row_data[4],
        "bcum": row_data[5],
        "bc_g": row_data[6],
        "design": row_data[7],
        "d_g": row_data[8],
        "p1": row_data[9],
        "p1_g": row_data[10],
        "e1": row_data[11],
        "e1_g": row_data[12],
        "wd": row_data[13],
        "wd_g": row_data[14],
        "algo": row_data[15],
        "al_g": row_data[16],
        "p2": row_data[17],
        "p2_g": row_data[18],
        "e2": row_data[19],
        "e2_g": row_data[20],
        "sd": row_data[21],
        "sd_g": row_data[22],
        "js": row_data[23],
        "js_g": row_data[24],
        "php": row_data[25],
        "ph_g": row_data[26],
        "db": row_data[27],
        "db_g": row_data[28],
        "vc1": row_data[29],
        "v1_g": row_data[30],
        "node": row_data[31],
        "no_g": row_data[32],
        "e3": row_data[33],
        "e3_g": row_data[34],
        "p3": row_data[35],
        "p3_g": row_data[36],
        "oop": row_data[37],
        "op_g": row_data[38],
        "lar": row_data[39],
        "la_g": row_data[40],
        "vue": row_data[41],
        "vu_g": row_data[42],
        "vc2": row_data[43],
        "v2_g": row_data[44],
        "e4": row_data[45],
        "e4_g": row_data[46],
        "p4": row_data[47],
        "p4_g": row_data[48],
        "int": row_data[49],
        "in_g": row_data[50],
        'cur_date': current_date
    })
    doc_name = os.path.join(output_directory, f"{row_data[1]}.docx")
    doc.save(doc_name)
    return doc_name

def TranscriptPdf(doc_path, pdf_directory):
    pdf_path = os.path.join(pdf_directory, os.path.splitext(os.path.basename(doc_path))[0] + ".pdf")
    convert(doc_path, pdf_path)
    return pdf_path

def generate_transcripts(option):
    excel_file = "data.xlsx"
    template_file = "template-pnc.docx"
    docx_directory = "Transcript_Doc"
    pdf_directory = "Transcript_PDF"

    os.makedirs(docx_directory, exist_ok=True)
    os.makedirs(pdf_directory, exist_ok=True)
    data_rows = TranscriptExcel_data(excel_file)

    for row in data_rows[1:]:
        if option in ["doc", "both"]:
            doc_path = TranscriptDocument(template_file, docx_directory, row)
        if option in ["pdf", "both"]:
            if option == "pdf":
                doc_path = TranscriptDocument(template_file, pdf_directory, row)
            TranscriptPdf(doc_path, pdf_directory)
            if option == "pdf":
                os.remove(doc_path)
    print(f"All files for option '{option}' have been generated!")

# GUI Implementation
def create_ui():
    def show_option_menu(title, generate_function):
        option_window = tk.Toplevel(window)
        option_window.title(title)
        tk.Label(option_window, text="Select the output format:").pack(padx=20, pady=10)

        def generate_with_option(selected_option):
            generate_function(selected_option, notify=lambda msg: messagebox.showinfo("Success", msg))
            option_window.destroy()

        tk.Button(option_window, text="DOCX Only", command=lambda: generate_with_option("doc")).pack(padx=20, pady=5)
        tk.Button(option_window, text="PDF Only", command=lambda: generate_with_option("pdf")).pack(padx=20, pady=5)
        tk.Button(option_window, text="Both DOCX and PDF", command=lambda: generate_with_option("both")).pack(padx=20, pady=5)

    window = tk.Tk()
    window.geometry("400x400")
    window.title("Automated Document Generation")

    tk.Label(window, text="Automated Document Generation", font=("Arial", 16, "bold"), fg="blue").pack(pady=20)

    tk.Button(window, text="Generate Transcripts", font=("Arial", 12), bg="blue", fg="white", width=20,
              command=lambda: show_option_menu("Generate Transcripts", generate_transcripts)).pack(pady=10)

    tk.Button(window, text="Generate Certificates", font=("Arial", 12), bg="blue", fg="white", width=20,
              command=lambda: generate_certificates(
                  excel_file="Certificate.xlsx",
                  template_file="certificate.png",
                  output_folder="Certificates",
                  notify=lambda msg: messagebox.showinfo("Success", msg)
              )).pack(pady=10)

    tk.Button(window, text="Generate Associate", font=("Arial", 12), bg="blue", fg="white", width=20,
              command=lambda: show_option_menu("Generate Associate", GeneratAssociate)).pack(pady=10)

    tk.Button(
        window, 
        text="Generate All", 
        font=("Arial", 12), 
        bg="green", 
        fg="white", 
        width=20,
        command=lambda: show_option_menu(
            "Generate All Documents", 
            lambda option: (
                generate_certificates(
                    excel_file="Certificate.xlsx",
                    template_file="certificate.png",
                    output_folder="Certificates",
                    notify=lambda msg: messagebox.showinfo("Success", msg)
                ),
                generate_transcripts(option, notify=lambda msg: messagebox.showinfo("Success", msg)),
                GeneratAssociate(option, notify=lambda msg: messagebox.showinfo("Success", msg))
            )
        )
    ).pack(pady=10)

    window.mainloop()

# Update your functions to accept a `notify` parameter
def generate_certificates(excel_file, template_file, output_folder, font_path="arialbd.ttf", font_size=100, notify=None):
    data = pd.read_excel(excel_file)
    if not os.path.exists(output_folder):
        os.makedirs(output_folder)
    font_name = ImageFont.truetype(font_path, font_size)
    for index, row in data.iterrows():
        name = row["Name"]
        certificate = Image.open(template_file)
        draw = ImageDraw.Draw(certificate)
        if len(name) >= 15 and len(name) < 25:
            name_position = (550, 600)
        elif len(name) >= 10 and len(name) < 15:
            name_position = (700, 600)
        else:
            name_position = (730, 600)
        draw.text(name_position, name, fill="orange", font=font_name)
        output_path = os.path.join(output_folder, "certificate_" + name + ".png")
        certificate.save(output_path)
        print(f"Certificate generated for {name} and saved to {output_path}")
    if notify:
        notify("All certificates have been generated!")

def GeneratAssociate(option, notify=None):
    excel_file = "Associate.xlsx"
    template_file = "WEP Temporary Certificate - Template.docx"
    docx_directory = "Associate_Documents"
    pdf_directory = "Associate_PDF"

    os.makedirs(docx_directory, exist_ok=True)
    os.makedirs(pdf_directory, exist_ok=True)
    data_rows = AssociateExcel_data(excel_file)

    for row in data_rows[1:]:
        if option in ["doc", "both"]:
            doc_path = AssociateDocument(template_file, docx_directory, row)
        if option in ["pdf", "both"]:
            if option == "pdf":
                doc_path = AssociateDocument(template_file, pdf_directory, row)
            AssociateConvertPDF(doc_path, pdf_directory)
            if option == "pdf":
                os.remove(doc_path)
    if notify:
        notify(f"All files for option '{option}' have been generated!")

def generate_transcripts(option, notify=None):
    excel_file = "data.xlsx"
    template_file = "template-pnc.docx"
    docx_directory = "Transcript_Doc"
    pdf_directory = "Transcript_PDF"

    os.makedirs(docx_directory, exist_ok=True)
    os.makedirs(pdf_directory, exist_ok=True)
    data_rows = TranscriptExcel_data(excel_file)

    for row in data_rows[1:]:
        if option in ["doc", "both"]:
            doc_path = TranscriptDocument(template_file, docx_directory, row)
        if option in ["pdf", "both"]:
            if option == "pdf":
                doc_path = TranscriptDocument(template_file, pdf_directory, row)
            TranscriptPdf(doc_path, pdf_directory)
            if option == "pdf":
                os.remove(doc_path)
    if notify:
        notify(f"All files for option '{option}' have been generated!")

# Run the Application
if __name__ == "__main__":
    create_ui()
